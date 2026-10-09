from __future__ import annotations

import csv
import io
import json
import re
import statistics
import subprocess
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from typing import Any

import pymupdf as fitz
from openpyxl import load_workbook
from PIL import Image


class ExtractionError(RuntimeError):
    pass


@dataclass(frozen=True)
class Region:
    region_id: str
    page_number: int
    region_type: str
    source_text: str
    normalized_polygon: tuple[tuple[float, float], ...]
    pixel_polygon: tuple[tuple[int, int], ...] | None
    confidence: float | None
    unresolved: bool
    style: dict[str, Any]


@dataclass(frozen=True)
class ExtractionResult:
    extraction_id: str
    method: str
    engine_name: str
    engine_version: str
    language_profile: str | None
    regions: tuple[Region, ...]
    unresolved_regions: int
    page_coverage: tuple[dict[str, Any], ...] = ()

    @property
    def document_complete(self) -> bool:
        return bool(self.page_coverage) and all(
            row["status"] == "NATIVE_TEXT_PRESENT" for row in self.page_coverage
        )

    def as_dict(self) -> dict[str, Any]:
        return {
            "extraction_id": self.extraction_id,
            "method": self.method,
            "engine_name": self.engine_name,
            "engine_version": self.engine_version,
            "language_profile": self.language_profile,
            "regions": [asdict(region) for region in self.regions],
            "unresolved_regions": self.unresolved_regions,
            "page_coverage": list(self.page_coverage),
            "document_complete": self.document_complete,
        }


@dataclass(frozen=True)
class Correction:
    correction_id: str
    extraction_id: str
    region_id: str
    corrected_text: str | None
    corrected_polygon: tuple[tuple[float, float], ...] | None
    reason: str
    corrected_by: str
    review_state: str
    created_at: str
    supersedes_correction_id: str | None = None


class CorrectionLedger:
    def __init__(self) -> None:
        self._corrections: list[Correction] = []

    def propose(
        self,
        result: ExtractionResult,
        region_id: str,
        *,
        corrected_text: str | None,
        corrected_polygon: tuple[tuple[float, float], ...] | None,
        reason: str,
        corrected_by: str,
    ) -> Correction:
        if not any(region.region_id == region_id for region in result.regions):
            raise ExtractionError("correction refers to an unknown immutable region")
        if not reason.strip() or not corrected_by.strip():
            raise ExtractionError("correction reason and actor are required")
        correction = Correction(
            str(uuid.uuid4()), result.extraction_id, region_id, corrected_text,
            corrected_polygon, reason, corrected_by, "proposed",
            datetime.now(timezone.utc).isoformat(),
        )
        self._corrections.append(correction)
        return correction

    def decide(self, correction_id: str, *, decision: str, reviewer_id: str, reason: str) -> Correction:
        if decision not in {"accepted", "rejected"}:
            raise ExtractionError("correction decision must be accepted or rejected")
        if not reviewer_id.strip() or not reason.strip():
            raise ExtractionError("reviewer and decision reason are required")
        proposal = next((item for item in self._corrections if item.correction_id == correction_id), None)
        if proposal is None or proposal.review_state != "proposed":
            raise ExtractionError("only a known proposed correction can be decided")
        decision_record = Correction(
            str(uuid.uuid4()), proposal.extraction_id, proposal.region_id, proposal.corrected_text,
            proposal.corrected_polygon, reason, reviewer_id, decision,
            datetime.now(timezone.utc).isoformat(), proposal.correction_id,
        )
        self._corrections.append(decision_record)
        return decision_record

    def effective_text(self, result: ExtractionResult, region_id: str) -> str:
        original = next((region for region in result.regions if region.region_id == region_id), None)
        if original is None:
            raise ExtractionError("region not found")
        accepted = [
            correction for correction in self._corrections
            if correction.extraction_id == result.extraction_id
            and correction.region_id == region_id
            and correction.review_state == "accepted"
            and correction.corrected_text is not None
        ]
        return accepted[-1].corrected_text if accepted else original.source_text


def _rectangle_polygon(x0: float, y0: float, x1: float, y1: float, width: float, height: float) -> tuple[tuple[float, float], ...]:
    if width <= 0 or height <= 0:
        raise ExtractionError("invalid page dimensions")
    return (
        (x0 / width, y0 / height), (x1 / width, y0 / height),
        (x1 / width, y1 / height), (x0 / width, y1 / height),
    )


def _classify_text(text: str, size: float, median_size: float, flags: int, y1: float, height: float) -> str:
    normalized = text.strip()
    if not normalized:
        return "text"
    if size <= median_size * 0.78 or y1 >= height * 0.93:
        return "footnote"
    if size >= median_size * 1.25 or flags & 16:
        return "heading"
    if re.search(r"(?:ingredients?|tərkib|состав)\s*:", normalized, re.IGNORECASE):
        return "ingredient_list"
    return "text"


class LayoutExtractor:
    def __init__(self, *, tesseract_languages: str = "aze+eng+rus", minimum_word_confidence: float = 0.70) -> None:
        self.languages = tesseract_languages
        self.minimum_word_confidence = minimum_word_confidence

    def extract_pdf(self, data: bytes, *, force_ocr: bool = False) -> ExtractionResult:
        """Per-page extraction. Never let text on one page mask an unread page.

        A raster image next to native text can contain additional writing, so it
        explicitly requires review. OCR confidence is not a measured CER/WER.
        """
        try:
            document = fitz.open(stream=data, filetype="pdf")
        except Exception as exc:
            raise ExtractionError("invalid PDF") from exc
        extraction_id = str(uuid.uuid4())
        output: list[Region] = []
        pages: list[dict[str, Any]] = []
        methods: set[str] = set()
        try:
            if document.page_count < 1:
                raise ExtractionError("PDF contains no pages")
            for page_number, page in enumerate(document, 1):
                page_regions: list[Region] = []
                page_data = page.get_text("dict", sort=True) if not force_ocr else {"blocks": []}
                blocks = page_data.get("blocks", [])
                spans = [
                    span for block in blocks if block.get("type") == 0
                    for line in block.get("lines", []) for span in line.get("spans", [])
                    if str(span.get("text", "")).strip()
                ]
                native_chars = sum(len(str(span.get("text", "")).strip()) for span in spans)
                raster_images = sum(block.get("type") == 1 for block in blocks)
                if native_chars >= 20 and not force_ocr:
                    sizes = [float(span.get("size", 0)) for span in spans if float(span.get("size", 0)) > 0]
                    median_size = statistics.median(sizes) if sizes else 10.0
                    for span in spans:
                        x0, y0, x1, y1 = (float(value) for value in span["bbox"])
                        text = str(span["text"])
                        flags = int(span.get("flags", 0))
                        page_regions.append(Region(
                            str(uuid.uuid4()), page_number,
                            _classify_text(text, float(span.get("size", 0)), median_size,
                                           flags, y1, page.rect.height),
                            text, _rectangle_polygon(x0, y0, x1, y1, page.rect.width, page.rect.height),
                            None, None, False, {
                                "font": span.get("font"), "font_size": span.get("size"),
                                "bold": bool(flags & 16), "italic": bool(flags & 2),
                                "color": span.get("color"), "source_method": "native_pdf",
                            },
                        ))
                    methods.add("native_pdf")
                    status = "EXTRACTED_WITH_REVIEW_REQUIRED" if raster_images else "NATIVE_TEXT_PRESENT"
                    method = "native_pdf"
                else:
                    pixmap = page.get_pixmap(dpi=300, alpha=False)
                    page_regions = self._tesseract_regions(pixmap.tobytes("png"), page_number)
                    methods.add("ocr")
                    status = "OCR_REVIEW_REQUIRED" if page_regions else "UNRESOLVED_EMPTY_OCR"
                    method = "ocr"
                output.extend(page_regions)
                pages.append({
                    "page_number": page_number, "method": method, "status": status,
                    "region_count": len(page_regions), "native_character_count": native_chars,
                    "raster_image_count": raster_images,
                    "review_required": status not in {"NATIVE_TEXT_PRESENT"},
                    "ocr_confidence_is_accuracy": False,
                })
        finally:
            document.close()
        method = "hybrid_pdf" if len(methods) > 1 else next(iter(methods))
        return ExtractionResult(
            extraction_id, method, "PyMuPDF+Tesseract" if "ocr" in methods else "PyMuPDF",
            f"PyMuPDF {fitz.VersionBind}; Tesseract {self.tesseract_version()}" if "ocr" in methods else fitz.VersionBind,
            self.languages if "ocr" in methods else None, tuple(output),
            sum(r.unresolved for r in output), tuple(pages),
        )

    def extract_image(self, data: bytes) -> ExtractionResult:
        try:
            with Image.open(io.BytesIO(data)) as image:
                image.verify()
        except Exception as exc:
            raise ExtractionError("invalid image") from exc
        regions = self._tesseract_regions(data, 1)
        return ExtractionResult(
            str(uuid.uuid4()), "ocr", "tesseract", self.tesseract_version(), self.languages,
            tuple(regions), sum(region.unresolved for region in regions),
            ({"page_number": 1, "method": "ocr",
              "status": "OCR_REVIEW_REQUIRED" if regions else "UNRESOLVED_EMPTY_OCR",
              "region_count": len(regions), "review_required": True,
              "ocr_confidence_is_accuracy": False},),
        )

    def _tesseract_regions(self, image: bytes, page_number: int) -> list[Region]:
        command = ["tesseract", "stdin", "stdout", "-l", self.languages, "--dpi", "300", "tsv"]
        try:
            completed = subprocess.run(
                command, input=image, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                timeout=120, check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise ExtractionError("OCR engine unavailable or timed out") from exc
        if completed.returncode != 0:
            message = completed.stderr.decode("utf-8", "replace")[:500]
            raise ExtractionError(f"OCR failed: {message}")
        with Image.open(io.BytesIO(image)) as decoded:
            width, height = decoded.size
        reader = csv.DictReader(io.StringIO(completed.stdout.decode("utf-8", "replace")), delimiter="\t")
        regions: list[Region] = []
        for row in reader:
            text = str(row.get("text", "")).strip()
            if not text:
                continue
            try:
                confidence = max(0.0, min(1.0, float(row["conf"]) / 100.0))
                left, top = int(row["left"]), int(row["top"])
                box_width, box_height = int(row["width"]), int(row["height"])
            except (KeyError, TypeError, ValueError):
                continue
            regions.append(Region(
                str(uuid.uuid4()), page_number, "word", text,
                _rectangle_polygon(left, top, left + box_width, top + box_height, width, height),
                ((left, top), (left + box_width, top), (left + box_width, top + box_height), (left, top + box_height)),
                confidence, confidence < self.minimum_word_confidence, {},
            ))
        return regions

    @staticmethod
    def tesseract_version() -> str:
        try:
            completed = subprocess.run(
                ["tesseract", "--version"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                timeout=10, check=False, text=True,
            )
            return completed.stdout.splitlines()[0].removeprefix("tesseract ").strip()
        except (OSError, subprocess.TimeoutExpired):
            return "unavailable"

    def extract_spreadsheet(self, data: bytes) -> ExtractionResult:
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except Exception as exc:
            raise ExtractionError("invalid spreadsheet") from exc
        regions: list[Region] = []
        for sheet_number, sheet in enumerate(workbook.worksheets, 1):
            max_row = max(1, sheet.max_row)
            max_column = max(1, sheet.max_column)
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = str(cell.value)
                    is_formula = value.startswith("=")
                    x0 = (cell.column - 1) / max_column
                    x1 = cell.column / max_column
                    y0 = (cell.row - 1) / max_row
                    y1 = cell.row / max_row
                    regions.append(Region(
                        str(uuid.uuid4()), sheet_number, "table_cell", value,
                        ((x0, y0), (x1, y0), (x1, y1), (x0, y1)), None,
                        None if is_formula else 1.0, is_formula,
                        {"sheet": sheet.title, "coordinate": cell.coordinate, "formula": is_formula},
                    ))
        workbook.close()
        return ExtractionResult(
            str(uuid.uuid4()), "spreadsheet", "openpyxl", "3", None, tuple(regions),
            sum(region.unresolved for region in regions),
        )

    def extract_csv(self, data: bytes) -> ExtractionResult:
        try:
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ExtractionError("invalid CSV") from exc
        column_count = max((len(row) for row in rows), default=1)
        row_count = max(len(rows), 1)
        regions = tuple(
            Region(
                str(uuid.uuid4()), 1, "table_cell", value,
                (((column - 1) / column_count, (row_index - 1) / row_count),
                 (column / column_count, (row_index - 1) / row_count),
                 (column / column_count, row_index / row_count),
                 ((column - 1) / column_count, row_index / row_count)),
                None, 1.0, False, {"row": row_index, "column": column},
            )
            for row_index, row in enumerate(rows, 1)
            for column, value in enumerate(row, 1)
        )
        return ExtractionResult(str(uuid.uuid4()), "spreadsheet", "python-csv", "1", None, regions, 0)

    def extract_structured_record(self, data: bytes) -> ExtractionResult:
        try:
            value = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError, RecursionError) as exc:
            raise ExtractionError("invalid structured API record") from exc
        regions: list[Region] = []

        def walk(node: Any, path: str, depth: int) -> None:
            if depth > 64:
                raise ExtractionError("structured record nesting exceeds policy")
            if isinstance(node, dict):
                for key in sorted(node):
                    walk(node[key], f"{path}/{key}", depth + 1)
            elif isinstance(node, list):
                for index, item in enumerate(node):
                    walk(item, f"{path}/{index}", depth + 1)
            else:
                regions.append(Region(
                    str(uuid.uuid4()), 1, "structured_field", json.dumps(node, ensure_ascii=False),
                    ((0.0, 0.0), (1.0, 0.0), (1.0, 1.0), (0.0, 1.0)), None, 1.0, False,
                    {"json_pointer": path or "/"},
                ))

        walk(value, "", 0)
        return ExtractionResult(
            str(uuid.uuid4()), "structured_api", "python-json", "1", None, tuple(regions), 0
        )


def label_declaration_candidates(result: ExtractionResult) -> dict[str, Any]:
    """Derive review-only ingredient/allergen statement candidates.

    This is deliberately lexical: it preserves region references and confidence
    and never promotes OCR text into an accepted label record automatically.
    """
    # Extractors already emit reading order; region IDs are random custody IDs
    # and must never be used as a layout sort key.
    ordered = list(result.regions)
    combined = "\n".join(region.source_text for region in ordered if region.source_text.strip())
    patterns = {
        "ingredient_declarations": re.compile(
            r"(?:ingredients?|tərkib(?:i)?|состав)\s*[:\-]\s*"
            r"(.+?)(?=\s+(?:allergens?|contains|allergenlər?|аллерген(?:ы)?)\s*[:\-]|\n|$)",
            re.IGNORECASE,
        ),
        "allergen_statements": re.compile(
            r"(?:allergens?|contains|allergenlər?|аллерген(?:ы)?)\s*[:\-]\s*([^\n]+)",
            re.IGNORECASE,
        ),
    }
    candidates: dict[str, list[dict[str, Any]]] = {key: [] for key in patterns}
    for key, pattern in patterns.items():
        for match in pattern.finditer(combined):
            candidates[key].append({
                "text": match.group(1).strip(),
                "matched_text": match.group(0).strip(),
                "character_span": [match.start(), match.end()],
                "review_required": True,
            })
    for region in ordered:
        if region.region_type == "ingredient_list" and not any(
            item["matched_text"] in region.source_text.strip()
            for item in candidates["ingredient_declarations"]
        ):
            candidates["ingredient_declarations"].append({
                "text": region.source_text.strip(), "matched_text": region.source_text.strip(),
                "region_id": region.region_id, "confidence": region.confidence,
                "review_required": True,
            })
    return {**candidates, "source_extraction_id": result.extraction_id,
            "unresolved_regions": result.unresolved_regions, "page_coverage": list(result.page_coverage),
            "document_complete": result.document_complete, "automatic_acceptance": False}


def _levenshtein(left: list[str], right: list[str]) -> int:
    previous = list(range(len(right) + 1))
    for left_index, left_item in enumerate(left, 1):
        current = [left_index]
        for right_index, right_item in enumerate(right, 1):
            current.append(min(
                current[-1] + 1,
                previous[right_index] + 1,
                previous[right_index - 1] + (left_item != right_item),
            ))
        previous = current
    return previous[-1]


@dataclass(frozen=True)
class BenchmarkReport:
    character_error_rate: float
    word_error_rate: float
    passed: bool
    thresholds: dict[str, float]


def benchmark_text(reference: str, hypothesis: str, *, max_cer: float, max_wer: float) -> BenchmarkReport:
    reference_characters = list(reference)
    hypothesis_characters = list(hypothesis)
    reference_words = reference.split()
    hypothesis_words = hypothesis.split()
    character_error_rate = _levenshtein(reference_characters, hypothesis_characters) / max(1, len(reference_characters))
    word_error_rate = _levenshtein(reference_words, hypothesis_words) / max(1, len(reference_words))
    return BenchmarkReport(
        character_error_rate, word_error_rate,
        character_error_rate <= max_cer and word_error_rate <= max_wer,
        {"max_character_error_rate": max_cer, "max_word_error_rate": max_wer},
    )
